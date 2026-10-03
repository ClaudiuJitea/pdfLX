"""pdf2docx adapter, native text export, and atomic output regression coverage."""
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import Mock, patch
import pymupdf as fitz
from pdflx import pdf_handler


class DocumentExportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        with fitz.open() as source, fitz.open() as target:
            source.new_page(width=300, height=400)
            source[0].insert_text((30, 50), 'Compressed form invoice')
            source[0].insert_text((30, 80), 'Amount 123.45')
            target.new_page(width=500, height=600)
            target[0].show_pdf_page(fitz.Rect(30, 40, 330, 440), source, 0)
            self.raw = target.tobytes(deflate=True)
        self.doc = fitz.open(stream=self.raw, filetype='pdf')
        self.addCleanup(self.doc.close)

    def converter(self, output=b'converted document', error=None):
        instance = Mock()
        def convert(path, **kwargs):
            Path(path).write_bytes(output)
            if error:
                raise error
        instance.convert.side_effect = convert
        factory = Mock(return_value=instance)
        return factory, instance

    def test_word_adapter_uses_snapshot_strict_settings_and_closes_converter(self):
        self.doc[0].insert_text((30, 500), 'Unsaved edit stays')
        before = self.doc[0].get_pixmap().samples
        factory, converter = self.converter()
        destination = self.root/'result.docx'
        with patch.object(pdf_handler,'HAS_PDF2DOCX',True), patch.object(pdf_handler,'Pdf2DocxConverter',factory):
            ok, error = pdf_handler.export_document(self.doc,output_path=destination)
        self.assertTrue(ok,error)
        data = factory.call_args.kwargs['stream']
        with fitz.open(stream=data,filetype='pdf') as payload:
            text = payload[0].get_text()
            self.assertIn('Compressed form invoice',text)
            self.assertIn('Unsaved edit stays',text)
        kwargs = converter.convert.call_args.kwargs
        self.assertEqual(kwargs,dict(multi_processing=False,ignore_page_error=False,raw_exceptions=True))
        converter.close.assert_called_once()
        self.assertEqual(destination.read_bytes(),b'converted document')
        self.assertEqual(self.doc[0].get_pixmap().samples,before)
        self.assertFalse(list(self.root.glob('.pdflx-export-*')))

    def test_native_text_export_reads_compressed_forms_and_current_edits(self):
        self.doc[0].insert_text((30,500),'Unsaved edit stays')
        destination=self.root/'result.txt'
        with patch.object(pdf_handler,'HAS_PDF2DOCX',False):
            ok,error=pdf_handler.export_document(self.doc,output_path=destination,target_format='txt')
        self.assertTrue(ok,error)
        text=destination.read_text()
        for value in ('Compressed form invoice','Amount 123.45','Unsaved edit stays'):
            self.assertIn(value,text)

    def test_password_protected_source_is_authenticated_before_conversion(self):
        path=self.root/'encrypted.pdf'
        self.doc.save(path,encryption=fitz.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='reader')
        destination=self.root/'result.docx'
        factory,converter=self.converter()
        with patch.object(pdf_handler,'HAS_PDF2DOCX',True),patch.object(pdf_handler,'Pdf2DocxConverter',factory):
            ok,error=pdf_handler.export_document(source_pdf_path=path,output_path=destination,password='reader')
            self.assertTrue(ok,error)
            with fitz.open(stream=factory.call_args.kwargs['stream'],filetype='pdf') as payload:
                self.assertFalse(payload.needs_pass)
                self.assertIn('Compressed form invoice',payload[0].get_text())
            for password in ('','wrong'):
                factory.reset_mock()
                ok,error=pdf_handler.export_document(source_pdf_path=path,output_path=destination,password=password)
                self.assertFalse(ok)
                self.assertIn('password',error.lower())
                factory.assert_not_called()
                self.assertEqual(destination.read_bytes(),b'converted document')

    def test_failed_page_empty_output_and_replace_failure_preserve_destination(self):
        destination=self.root/'existing.docx'
        destination.write_bytes(b'previous export')
        for output,error,replace_error in ((b'partial',ValueError('Page 2 failed'),None),(b'',None,None),
                                          (b'complete',None,OSError('Destination unavailable'))):
            factory,converter=self.converter(output,error)
            with patch.object(pdf_handler,'HAS_PDF2DOCX',True),patch.object(pdf_handler,'Pdf2DocxConverter',factory):
                if replace_error:
                    with patch.object(pdf_handler.os,'replace',side_effect=replace_error):
                        ok,message=pdf_handler.export_document(self.doc,output_path=destination)
                else:
                    ok,message=pdf_handler.export_document(self.doc,output_path=destination)
            self.assertFalse(ok)
            self.assertTrue(message)
            converter.close.assert_called_once()
            self.assertEqual(destination.read_bytes(),b'previous export')
            self.assertFalse(list(self.root.glob('.pdflx-export-*')))

    def test_missing_engine_and_unsupported_formats_are_clear(self):
        with patch.object(pdf_handler,'HAS_PDF2DOCX',False):
            ok,error=pdf_handler.export_document(self.doc,output_path=self.root/'out.docx')
        self.assertFalse(ok)
        self.assertIn('pdf2docx',error)
        for fmt in ('pptx','odt','odp'):
            ok,error=pdf_handler.export_document(self.doc,output_path=self.root/'out',target_format=fmt)
            self.assertFalse(ok)
            self.assertIn('Supported formats are DOCX and TXT',error)
        from pdflx.export_dialog import FORMAT_ITEMS
        self.assertEqual([item[0] for item in FORMAT_ITEMS],['DOCX','TXT'])

    @unittest.skipUnless(pdf_handler.HAS_PDF2DOCX,'pdf2docx is not installed')
    def test_real_pdf2docx_conversion_preserves_text_and_tables(self):
        from pdflx.table_creation import create_table_objects
        from pdflx.models import EditableText,EditableShape
        page=self.doc[0]
        pdf_handler.save_page_snapshot(self.doc,0,force=True)
        objects=create_table_objects(page,[['Item','Value'],['Alpha','0012']],font_size=10)
        texts=[o for o in objects if isinstance(o,EditableText)]
        shapes=[o for o in objects if isinstance(o,EditableShape)]
        ok,error=pdf_handler.rebuild_page(self.doc,0,texts,shapes,[])
        self.assertTrue(ok,error)
        destination=self.root/'table.docx'
        ok,error=pdf_handler.export_document(self.doc,output_path=destination)
        self.assertTrue(ok,error)
        from docx import Document
        word=Document(destination)
        self.assertTrue(word.tables)
        contents=[[cell.text for cell in row.cells] for table in word.tables for row in table.rows]
        self.assertIn(['Alpha','0012'],contents)
        with zipfile.ZipFile(destination) as archive:
            self.assertIsNone(archive.testzip())
